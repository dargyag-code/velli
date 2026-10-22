# Uso, fuera del repo (pycel no es dependencia del proyecto):
#   python -m venv venv && venv/Scripts/pip install pycel openpyxl
#   venv/Scripts/python verificar_plantilla.py <plantilla-precios.xlsx> <copia-de-prueba.xlsx>
# Salida guardada de la corrida del 2026-10-22: verificar_plantilla_salida.txt
#
# Verifica la plantilla con pycel (evaluador de fórmulas de Excel en Python).
# pycel evalúa las dos ramas de IF, a diferencia de Excel/LibreOffice/Google
# Sheets; por eso se evalúan las filas CON datos y se comprueba que las filas
# vacías del Registro tienen exactamente la misma fórmula (salvo el número
# de fila). Escenario 1: el ejemplo del archivo. Escenario 2: una copia con
# tres casos más (corte sin fórmula, otra estilista, diferencias fuera de 5 %).
# Escenario 3: los avisos del estándar con fibra sensibilizada.
import re
import sys
from datetime import date
from openpyxl import load_workbook
from pycel import ExcelCompiler

ruta, copia = sys.argv[1], sys.argv[2]
copia_avisos = re.sub(r'\.xlsx$', '', copia) + '-avisos.xlsx'
fallas = 0

def revisar(xl, esperado, titulo):
    global fallas
    print(f'\n== {titulo}')
    for ref, valor in esperado.items():
        obtenido = xl.evaluate(ref)
        if isinstance(valor, float) or isinstance(valor, int):
            ok = isinstance(obtenido, (int, float)) and abs(obtenido - valor) < 1e-4
        else:
            ok = obtenido == valor
        fallas += 0 if ok else 1
        print(('OK   ' if ok else 'FALLA'), ref, '· esperado', valor, '· obtenido', obtenido)

def errores_en_filas_con_datos(ruta_xlsx, xl):
    global fallas
    wb = load_workbook(ruta_xlsx)
    total, errs = 0, []
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if not (isinstance(c.value, str) and c.value.startswith('=')):
                    continue
                # Filas de tablas sin datos: Excel devuelve "" por el IF; pycel no sirve ahí.
                if ws.title == 'Registro' and ws.cell(row=c.row, column=4).value in (None, ''):
                    continue
                if ws.title == 'Productos' and ws.cell(row=c.row, column=1).value in (None, ''):
                    continue
                total += 1
                try:
                    v = xl.evaluate(f"'{ws.title}'!{c.coordinate}")
                    if isinstance(v, str) and v.startswith('#'):
                        errs.append((ws.title, c.coordinate, v))
                except Exception as e:  # noqa: BLE001
                    errs.append((ws.title, c.coordinate, type(e).__name__))
    fallas += len(errs)
    print(f'fórmulas evaluadas en filas con datos: {total} · con error: {len(errs)}')
    for e in errs[:10]:
        print('  ERROR', e)

# ── Escenario 1: el archivo tal como se entrega ─────────────────────────────
# Gramos = 60 × 1,5 (medio) × 1,3 (alta) × 1,15 (canas 50–75 %) = 134,55.
xl = ExcelCompiler(filename=ruta)
errores_en_filas_con_datos(ruta, xl)
gramos1 = 60 * 1.5 * 1.3 * 1.15
costo1 = gramos1 * 500 + gramos1 * 1.5 * 30 + 3000
revisar(xl, {
    'Simulador!B15': 120000, 'Simulador!B16': 15000, 'Simulador!B17': 0.35, 'Simulador!B19': 182000,
    'Simulador!B21': '', 'Simulador!B27': 1.15, 'Simulador!B28': gramos1, 'Simulador!B29': gramos1 * 500,
    'Simulador!B30': gramos1 * 1.5, 'Simulador!B31': gramos1 * 1.5 * 30, 'Simulador!B32': 3000,
    'Simulador!B33': costo1, 'Simulador!B36': 0.40, 'Simulador!B38': 72800,
    'Simulador!B39': 182000 - costo1 - 72800, 'Simulador!B40': (182000 - costo1 - 72800) / 182000,
    'Registro!I5': 182000, 'Registro!K5': 8000, 'Registro!L5': 190000 / 182000 - 1, 'Registro!P5': 1,
    'Resumen!B4': 1, 'Resumen!B5': 1, 'Resumen!B6': 190000 / 182000 - 1,
    'Resumen!A9': 'Ejemplo · Caro', 'Resumen!B9': 1, 'Resumen!D9': 0, 'Resumen!A10': '',
    'Productos!G5': 500, 'Productos!G6': 30, 'Consumo!B7': 'Corto', 'Consumo!B13': 'Alta',
    'Consumo!B14': '0–25 %', 'Consumo!B17': '75–100 %', 'Consumo!C16': 1.15, 'Consumo!C17': 1.2,
}, 'Escenario 1: ejemplo del documento de diseño (precio 182.000, margen 32.870)')

# ── Escenario 2: copia con más casos ────────────────────────────────────────
wb = load_workbook(ruta)
s = wb['Servicios']
for col, v in zip('BCDEFGHIJ', ['corte', 40000, 'Sí', 'No', 'No', 'No', 'No', None, 0]):
    s[f'{col}10'] = v           # Corte
for col, v in zip('BCDEFGHIJ', ['color', 150000, 'Sí', 'Sí', 'Sí', 'Sí', 'Sí', 90, 5000]):
    s[f'{col}11'] = v           # Color y mechas
c = wb['Comisiones']
for col, v in zip('ABCD', ['Dani', 0.35, 0.30, 'Precio menos producto']):
    c[f'{col}6'] = v
r = wb['Registro']
casos = [
    (6, ['Dani', 'A. P.', 'Corte', 'Extra largo', 'Baja', 'Sensibilizada / muy dañada', '0–25 %'], 60000),
    (7, ['Ejemplo · Caro', 'M. R.', 'Color y mechas', 'Largo', 'Media', 'Procesada / porosa', '75–100 %'], 259000),
    (8, ['Dani', 'S. T.', 'Ejemplo · Tinte completo', 'Corto', 'Media', 'Sana / virgen', '0–25 %'], 130000),
]
for fila, datos, cobrado in casos:
    r[f'A{fila}'] = date.today()
    for col, v in zip('BCDEFGH', datos):
        r[f'{col}{fila}'] = v
    r[f'J{fila}'] = cobrado
sim = wb['Simulador']
for celda, v in {'B4': 'Color y mechas', 'B5': 'Largo', 'B6': 'Media', 'B7': 'Procesada / porosa',
                 'B8': '75–100 %', 'B9': 'Dani'}.items():
    sim[celda] = v
wb.save(copia)

# Gramos = 90 × 2,0 (largo) × 1,0 (media) × 1,2 (canas 75–100 %) = 216.
xl2 = ExcelCompiler(filename=copia)
errores_en_filas_con_datos(copia, xl2)
gramos2 = 90 * 2.0 * 1.0 * 1.2
costo2 = gramos2 * 500 + gramos2 * 1.5 * 30 + 5000
revisar(xl2, {
    'Registro!I6': 64000, 'Registro!L6': 60000 / 64000 - 1, 'Registro!P6': 0,
    'Registro!I7': 259000, 'Registro!K7': 0, 'Registro!P7': 1,
    'Registro!I8': 120000, 'Registro!K8': 10000, 'Registro!P8': 0,
    'Simulador!B19': 259000, 'Simulador!B21': 'Sugerir un tratamiento sellante. Técnica de prepigmentación.',
    'Simulador!B27': 1.2, 'Simulador!B28': gramos2, 'Simulador!B29': gramos2 * 500,
    'Simulador!B31': gramos2 * 1.5 * 30, 'Simulador!B33': costo2, 'Simulador!B36': 0.30,
    'Simulador!B37': 'Precio menos producto',
    'Simulador!B38': 0.30 * (259000 - costo2), 'Simulador!B39': 259000 - costo2 - 0.30 * (259000 - costo2),
    'Resumen!B4': 4, 'Resumen!B5': 0.5, 'Resumen!B6': 639000 / 625000 - 1,
    'Resumen!B9': 2, 'Resumen!C9': 449000 / 441000 - 1, 'Resumen!D9': 0,
    'Resumen!A10': 'Dani', 'Resumen!B10': 2, 'Resumen!C10': 190000 / 184000 - 1, 'Resumen!D10': 2,
}, 'Escenario 2: corte sin fórmula, otra estilista con % de color sobre precio menos producto')

# ── Escenario 3: fibra sensibilizada en un servicio químico ─────────────────
# El tinte de ejemplo no ajusta el precio por estado de la fibra: el precio
# no cambia, pero el aviso obligatorio aparece.
wb = load_workbook(ruta)
wb['Simulador']['B7'] = 'Sensibilizada / muy dañada'
wb.save(copia_avisos)
xl3 = ExcelCompiler(filename=copia_avisos)
revisar(xl3, {
    'Simulador!B19': 182000,
    'Simulador!B21': 'Obligatorio: prueba de mechón y aditivo protector Plex.',
}, 'Escenario 3: aviso obligatorio con fibra sensibilizada')

# ── Las 100 filas del Registro tienen la misma fórmula ──────────────────────
wb = load_workbook(ruta)
r = wb['Registro']
def normal(f, fila):
    # Fila relativa = dígitos pegados a la letra de la columna (A5, $O5); las
    # absolutas ($B$9) llevan $ antes del número y no cambian entre filas.
    return re.sub(rf'(?<=[A-Z]){fila}(?![0-9])', '#', f)
inconsistentes = 0
for col in 'IKLOP':
    base = normal(r[f'{col}5'].value, 5)
    for fila in range(6, 105):
        if normal(r[f'{col}{fila}'].value, fila) != base:
            inconsistentes += 1
fallas += inconsistentes
print(f'\nfórmulas del Registro distintas a la fila 5: {inconsistentes} (de {5 * 99})')
print('\nRESULTADO:', 'OK' if fallas == 0 else f'{fallas} FALLAS')
