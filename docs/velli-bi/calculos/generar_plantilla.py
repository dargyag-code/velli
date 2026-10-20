# ══════════════════════════════════════════════════════════════════════════
# Genera docs/velli-bi/plantilla-precios.xlsx: la plantilla para probar en
# papel el modelo de precios del cotizador con 3 salones (Etapa 3).
#
# Mismo modelo que docs/velli-bi/02-diseno-cotizador.md:
#   precio = redondeo((base + Σ suma) × (1 + Σ porcentaje))
#   costo  = gramos base × factor largo × factor densidad × $/g tinte
#            + ml oxidante × $/ml + insumos fijos
# Los valores de ejemplo reproducen el ejemplo del documento (§3): precio
# $182.000, margen $42.435.
#
# Uso: python docs/velli-bi/calculos/generar_plantilla.py <salida.xlsx>
# ══════════════════════════════════════════════════════════════════════════
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = sys.argv[1]

# ── Estilo (fuente Arial; convención: azul+amarillo = llenar, negro =
# fórmula, verde = dato traído de otra hoja) ──────────────────────────────
VERDE_VELLI = '2D5A27'
F_BASE = Font(name='Arial', size=10)
F_TITULO = Font(name='Arial', size=14, bold=True, color=VERDE_VELLI)
F_SUBTITULO = Font(name='Arial', size=11, bold=True, color=VERDE_VELLI)
F_NOTA = Font(name='Arial', size=9, italic=True, color='595959')
F_ENCABEZADO = Font(name='Arial', size=10, bold=True, color='FFFFFF')
F_INPUT = Font(name='Arial', size=10, color='0000FF')
F_INPUT_EJEMPLO = Font(name='Arial', size=10, color='0000FF', italic=True)
F_FORMULA = Font(name='Arial', size=10, color='000000')
F_FORMULA_FUERTE = Font(name='Arial', size=11, bold=True, color='000000')
F_ENLACE = Font(name='Arial', size=10, color='008000')
F_ETIQUETA = Font(name='Arial', size=10, bold=True)

FILL_INPUT = PatternFill('solid', start_color='FFFF00')
FILL_ENCABEZADO = PatternFill('solid', start_color=VERDE_VELLI)
FILL_RESULTADO = PatternFill('solid', start_color='E2EFDA')

BORDE = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
COP = '$#,##0;-$#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
FACTOR = '0.00'
GRAMOS = '#,##0.0'
FECHA = 'dd/mm/yyyy'

wb = Workbook()


def hoja(nombre, titulo, nota, anchos):
    ws = wb.create_sheet(nombre)
    ws['A1'] = titulo
    ws['A1'].font = F_TITULO
    if nota:
        ws['A2'] = nota
        ws['A2'].font = F_NOTA
        ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, len(anchos)))
        ws.row_dimensions[2].height = 30
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = ancho
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def encabezado(ws, fila, textos):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = F_ENCABEZADO
        c.fill = FILL_ENCABEZADO
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = BORDE
    ws.row_dimensions[fila].height = 32
    ws.print_title_rows = f'{fila}:{fila}'


def entrada(c, valor=None, fmt=None, ejemplo=False):
    if valor is not None:
        c.value = valor
    c.font = F_INPUT_EJEMPLO if ejemplo else F_INPUT
    c.fill = FILL_INPUT
    c.border = BORDE
    if fmt:
        c.number_format = fmt


def formula(c, f, fmt=None, enlace=False, fuerte=False):
    c.value = f
    c.font = F_ENLACE if enlace else (F_FORMULA_FUERTE if fuerte else F_FORMULA)
    c.border = BORDE
    if fmt:
        c.number_format = fmt


def lista(ws, fuente, rango, mensaje=None):
    # En el XML de .xlsx la fórmula de la lista va SIN '=' (así la guarda
    # Excel); openpyxl la copia tal cual, así que se quita aquí.
    dv = DataValidation(type='list', formula1=fuente.lstrip('='), allow_blank=True)
    dv.error = 'Elige un valor de la lista.'
    dv.errorTitle = 'Valor no válido'
    if mensaje:
        dv.prompt = mensaje
    ws.add_data_validation(dv)
    dv.add(rango)


# ══════════════════════════════════════════════════════════════════════════
# Cómo usarla
# ══════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = 'Cómo usarla'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 110
ws['A1'] = 'Velli · Plantilla de precios para probar en papel'
ws['A1'].font = F_TITULO
filas = [
    ('sub', 'Para qué sirve'),
    ('txt', 'Antes de construir el cotizador en la app, queremos comprobar con 3 salones reales dos cosas: '
            'si cuatro variables del diagnóstico (largo, densidad, estado de la fibra y canas) explican cómo cobran, '
            'y cuánto cambia lo cobrado según quién atiende.'),
    ('sub', 'Cómo usarla'),
    ('txt', '1. Haz una copia del archivo por salón (por ejemplo: plantilla-precios-SalonA.xlsx).'),
    ('txt', '2. Con el dueño, en 30 a 45 minutos: llena Servicios, Ajustes, Productos, Consumo y Comisiones.'),
    ('txt', '3. En el Simulador, prueba 3 o 4 casos reales y ajusta hasta que el dueño diga «sí, así cobro».'),
    ('txt', '4. Durante una semana, anota en Registro cada servicio real: la hoja calcula su precio y tú escribes lo que se cobró.'),
    ('txt', '5. Al final, mira el Resumen: cuántos casos coinciden con la hoja (±5 %) y cuánto se separa cada estilista.'),
    ('sub', 'Leyenda'),
    ('leyenda_input', 'Celda para llenar. Lo que ya trae en azul es un ejemplo o un valor de partida: cámbialo con el dueño.'),
    ('leyenda_formula', 'Fórmula: no la toques.'),
    ('leyenda_enlace', 'Dato traído de otra hoja.'),
    ('txt', 'Las filas «Ejemplo · …» muestran el formato esperado. Puedes sobrescribirlas.'),
    ('sub', 'Qué queremos aprender'),
    ('txt', '• ¿Bastan las cuatro variables, o el dueño cobra por algo más? Anótalo en Registro, columna «¿Algo que la hoja no considera?».'),
    ('txt', '• ¿El dueño piensa los ajustes como porcentaje, como suma fija o ambos?'),
    ('txt', '• ¿Cuánto se separa cada estilista del precio de la hoja?'),
    ('txt', '• ¿Los gramos que gasta el salón en color se parecen a los factores de la hoja Consumo?'),
    ('sub', 'Privacidad'),
    ('txt', 'En Registro escribe solo las iniciales de la clienta. No anotes datos de salud capilar ni del diagnóstico.'),
    ('sub', 'Fuente del modelo'),
    ('txt', 'docs/velli-bi/02-diseno-cotizador.md (secciones 1.1, 1.2 y 3). Precio = redondeo((base + sumas) × (1 + porcentajes)).'),
]
r = 3
for tipo, texto in filas:
    c = ws.cell(row=r, column=2, value=texto)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if tipo == 'sub':
        c.font = F_SUBTITULO
        r += 1
        continue
    c.font = F_BASE
    if tipo.startswith('leyenda'):
        muestra = ws.cell(row=r, column=1)
        muestra.border = BORDE
        if tipo == 'leyenda_input':
            muestra.value = 'Abc'
            muestra.font = F_INPUT
            muestra.fill = FILL_INPUT
        elif tipo == 'leyenda_formula':
            muestra.value = '123'
            muestra.font = F_FORMULA
        else:
            muestra.value = '123'
            muestra.font = F_ENLACE
    ws.row_dimensions[r].height = 28 if len(texto) > 110 else 16
    r += 1

# ══════════════════════════════════════════════════════════════════════════
# Servicios
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Servicios', 'Servicios del salón',
    'Precio base = lo que cobra el salón con cabello corto, densidad media, fibra sana y sin canas. '
    'Las columnas «Ajusta por…» dicen qué variables cambian el precio de ese servicio.',
    [30, 16, 14, 11, 11, 11, 11, 12, 12, 13, 11, 34],
)
for etiqueta, celda, valor, fmt in [
    ('Salón', 'B4', None, None), ('Ciudad', 'E4', None, None),
    ('Fecha de la visita', 'B5', None, FECHA), ('Redondeo del precio ($)', 'E5', 1000, COP),
    ('Vigencia de la cotización (días)', 'B6', 15, '0'),
]:
    col_etiqueta = 'A' if celda.startswith('B') else 'D'
    ws[f'{col_etiqueta}{celda[1:]}'] = etiqueta
    ws[f'{col_etiqueta}{celda[1:]}'].font = F_ETIQUETA
    entrada(ws[celda], valor, fmt)
ws['E5'].comment = Comment('Precio final redondeado a este múltiplo. 1000 = al millar más cercano.', 'Velli')
ws['B6'].comment = Comment('Valor de partida propuesto por Velli (15 días). Cámbialo si el salón usa otro.', 'Velli')

encabezado(ws, 8, [
    'Servicio', 'Categoría', 'Precio base ($)', 'Ajusta por largo', 'Ajusta por densidad',
    'Ajusta por estado de la fibra', 'Ajusta por canas', '¿Usa fórmula de color?',
    'Gramos de tinte base', 'Insumos fijos ($)', 'Duración (min)', 'Notas',
])
SERV_INI, SERV_FIN = 9, 38
ejemplo = ['Ejemplo · Tinte completo', 'color', 120000, 'Sí', 'Sí', 'No', 'Sí', 'Sí', 60, 3000, 120,
           'Precio de ejemplo para cabello corto y densidad media']
for col, valor in enumerate(ejemplo, start=1):
    entrada(ws.cell(row=SERV_INI, column=col), valor, None, ejemplo=True)
# Servicios sugeridos en el onboarding de la app (app/onboarding/page.tsx).
sugeridos = [
    ('Corte', 'corte'), ('Color y mechas', 'color'), ('Tratamiento capilar', 'tratamiento'),
    ('Keratina y alisado', 'alisado'), ('Rizos y definición', 'rizos'), ('Peinado y brushing', 'peinado'),
    ('Cuero cabelludo', 'cuero cabelludo'), ('Extensiones', 'extensiones'),
]
for i in range(SERV_INI + 1, SERV_FIN + 1):
    for col in range(1, 13):
        entrada(ws.cell(row=i, column=col))
    j = i - SERV_INI - 1
    if j < len(sugeridos):
        ws.cell(row=i, column=1, value=sugeridos[j][0])
        ws.cell(row=i, column=2, value=sugeridos[j][1])
for i in range(SERV_INI, SERV_FIN + 1):
    ws.cell(row=i, column=3).number_format = COP
    ws.cell(row=i, column=9).number_format = GRAMOS
    ws.cell(row=i, column=10).number_format = COP
ws.cell(row=SERV_INI + 1, column=1).comment = Comment(
    'Los 8 servicios que sugiere el onboarding de Velli. Renómbralos o bórralos según el salón.', 'Velli')
lista(ws, '"color,decoloración,tratamiento,corte,alisado,rizos,peinado,cuero cabelludo,extensiones,otro"',
      f'B{SERV_INI}:B{SERV_FIN}')
lista(ws, '"Sí,No"', f'D{SERV_INI}:H{SERV_FIN}')
ws.freeze_panes = f'B{SERV_INI}'
SERV = lambda col: f'Servicios!${col}${SERV_INI}:${col}${SERV_FIN}'

# ══════════════════════════════════════════════════════════════════════════
# Ajustes
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Ajustes', 'Ajustes por variable',
    'Cuánto sube el precio según cada variable. Puedes usar suma fija, porcentaje o ambos: '
    'precio = (base + sumas) × (1 + porcentajes). El valor de referencia (corto, media, sano, sin canas) siempre es 0.',
    [20, 16, 58, 14, 14],
)
encabezado(ws, 4, ['Variable', 'Valor', 'Cómo reconocerlo', 'Suma ($)', 'Porcentaje (%)'])
AJ_INI = 5
ajustes = [
    ('Largo', 'Corto', 'Hasta el mentón (referencia)', 0, 0, True),
    ('Largo', 'Medio', 'Hasta los hombros', 0, 0.20, False),
    ('Largo', 'Largo', 'Hasta los omóplatos', 0, 0.40, False),
    ('Largo', 'Extra largo', 'Hasta la cintura o más', 0, 0.60, False),
    ('Densidad', 'Baja', 'Se ve el cuero cabelludo a través del cabello', 0, 0, False),
    ('Densidad', 'Media', 'Se ve en parte (referencia)', 0, 0, True),
    ('Densidad', 'Alta', 'No se ve el cuero cabelludo', 0, 0.15, False),
    ('Estado de la fibra', 'Sano', 'Puntas sanas, sin daño, buena elasticidad (referencia)', 0, 0, True),
    ('Estado de la fibra', 'Sensibilizado', 'Puntas abiertas leves, daño mecánico o térmico, o elasticidad media', 10000, 0, False),
    ('Estado de la fibra', 'Comprometido', 'Puntas abiertas severas, daño químico o elasticidad baja', 20000, 0, False),
    ('Canas', 'Sin canas', 'Referencia', 0, 0, True),
    ('Canas', 'Menos del 25 %', 'Canas sueltas', 0, 0, False),
    ('Canas', '25–50 %', 'Hasta la mitad', 10000, 0, False),
    ('Canas', '50–75 %', 'Más de la mitad', 15000, 0, False),
    ('Canas', 'Más del 75 %', 'Casi todo cano', 25000, 0, False),
]
for k, (var, val, guia, suma, pct, ref) in enumerate(ajustes):
    fila = AJ_INI + k
    for col, v in [(1, var), (2, val), (3, guia)]:
        c = ws.cell(row=fila, column=col, value=v)
        c.font = F_BASE
        c.border = BORDE
    if ref:
        for col, v, fmt in [(4, 0, COP), (5, 0, PCT)]:
            c = ws.cell(row=fila, column=col, value=v)
            c.font = F_FORMULA
            c.number_format = fmt
            c.border = BORDE
    else:
        entrada(ws.cell(row=fila, column=4), suma, COP)
        entrada(ws.cell(row=fila, column=5), pct, PCT)
AJ_FIN = AJ_INI + len(ajustes) - 1
ws['D6'].comment = Comment('Valores de ejemplo, no recomendados: reemplázalos con los del salón.', 'Velli')
ws['C12'].comment = Comment('Regla propuesta por Velli para clasificar la fibra; la valida Keila (docs/velli-bi/02, §1.1).', 'Velli')
ws.freeze_panes = 'A5'
RANGO_LARGO = f'Ajustes!$B${AJ_INI}:$B${AJ_INI + 3}'
RANGO_DENS = f'Ajustes!$B${AJ_INI + 4}:$B${AJ_INI + 6}'
RANGO_ESTADO = f'Ajustes!$B${AJ_INI + 7}:$B${AJ_INI + 9}'
RANGO_CANAS = f'Ajustes!$B${AJ_INI + 10}:$B${AJ_INI + 14}'
AJ = lambda col: f'Ajustes!${col}${AJ_INI}:${col}${AJ_FIN}'

# ══════════════════════════════════════════════════════════════════════════
# Productos
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Productos', 'Productos y precio de compra',
    'Lo que más usa el salón y cuánto le cuesta. El costo por gramo o mililitro se calcula solo.',
    [32, 18, 14, 14, 10, 18, 18, 30],
)
encabezado(ws, 4, ['Producto', 'Marca', 'Tipo', 'Presentación', 'Unidad', 'Precio de compra ($)',
                   'Costo por g o ml ($)', 'Notas'])
PROD_INI, PROD_FIN = 5, 34
ejemplos_prod = [
    ['Ejemplo · Tinte 60 g', 'Marca X', 'Tinte', 60, 'g', 30000, 'Precio ficticio'],
    ['Ejemplo · Oxidante 20 vol', 'Marca X', 'Oxidante', 1000, 'ml', 30000, 'Precio ficticio'],
]
for i in range(PROD_INI, PROD_FIN + 1):
    datos = ejemplos_prod[i - PROD_INI] if i - PROD_INI < len(ejemplos_prod) else [None] * 7
    es_ejemplo = datos[0] is not None
    for col, v in zip([1, 2, 3, 4, 5, 6, 8], datos):
        entrada(ws.cell(row=i, column=col), v, None, ejemplo=es_ejemplo)
    ws.cell(row=i, column=6).number_format = COP
    formula(ws.cell(row=i, column=7), f'=IF(OR(D{i}="",F{i}="",D{i}=0),"",F{i}/D{i})', COP)
lista(ws, '"Tinte,Oxidante,Decolorante,Matizador,Tratamiento,Otro"', f'C{PROD_INI}:C{PROD_FIN}')
lista(ws, '"g,ml,unidad"', f'E{PROD_INI}:E{PROD_FIN}')
ws.freeze_panes = f'B{PROD_INI}'
PROD = lambda col: f'Productos!${col}${PROD_INI}:${col}${PROD_FIN}'

# ══════════════════════════════════════════════════════════════════════════
# Consumo
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Consumo', 'Consumo de producto en servicios de color',
    'La calculadora de color de Velli formula un tubo (60 g por defecto). Aquí el salón dice cuánto más gasta '
    'según el largo y la densidad. Pregúntale a la estilista que más color hace.',
    [44, 16, 18],
)
ws['A4'] = 'Oxidante por cada gramo de tinte (1:1,5 → 1,5)'
ws['A4'].font = F_ETIQUETA
entrada(ws['B4'], 1.5, FACTOR)
ws['B4'].comment = Comment('Valor de partida (orientativo, igual al de la calculadora sin carta de marca). Confírmalo en la carta técnica.', 'Velli')
encabezado(ws, 6, ['Variable', 'Valor', 'Factor de consumo'])
factores = [('Largo', 0, 1.0), ('Largo', 1, 1.5), ('Largo', 2, 2.0), ('Largo', 3, 3.0),
            ('Densidad', 4, 0.8), ('Densidad', 5, 1.0), ('Densidad', 6, 1.3)]
CONS_INI = 7
for k, (var, desfase, factor) in enumerate(factores):
    fila = CONS_INI + k
    c = ws.cell(row=fila, column=1, value=var)
    c.font = F_BASE
    c.border = BORDE
    formula(ws.cell(row=fila, column=2), f'=Ajustes!B{AJ_INI + desfase}', enlace=True)
    entrada(ws.cell(row=fila, column=3), factor, FACTOR)
CONS_FIN = CONS_INI + len(factores) - 1
ws['C7'].comment = Comment('Factores de partida propuestos por Velli, no medidos. Ajústalos con el salón: 2 = el doble de tinte que un cabello corto.', 'Velli')
CONS = lambda col: f'Consumo!${col}${CONS_INI}:${col}${CONS_FIN}'

# ══════════════════════════════════════════════════════════════════════════
# Comisiones
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Comisiones', 'Comisiones por estilista',
    'Porcentaje que recibe cada estilista. Si en color recibe otro porcentaje, llena esa columna; si no, se usa el general.',
    [30, 14, 22, 24, 34],
)
encabezado(ws, 4, ['Estilista', '% general', '% en servicios de color (opcional)', 'Se calcula sobre', 'Notas'])
COM_INI, COM_FIN = 5, 19
for i in range(COM_INI, COM_FIN + 1):
    ejemplo_fila = i == COM_INI
    valores = ['Ejemplo · Caro', 0.40, None, 'Precio', 'Comisión de ejemplo'] if ejemplo_fila else [None] * 5
    for col, v in enumerate(valores, start=1):
        entrada(ws.cell(row=i, column=col), v, None, ejemplo=ejemplo_fila)
    ws.cell(row=i, column=2).number_format = PCT
    ws.cell(row=i, column=3).number_format = PCT
lista(ws, '"Precio,Precio menos producto"', f'D{COM_INI}:D{COM_FIN}')
COM = lambda col: f'Comisiones!${col}${COM_INI}:${col}${COM_FIN}'
RANGO_ESTILISTAS = COM('A')


def suma_variables(col_aj, fila_serv, largo, dens, estado, canas):
    """Σ de la columna de Ajustes (D = suma, E = porcentaje) para las variables que aplican al servicio."""
    partes = []
    for col_serv, variable, valor in [('D', 'Largo', largo), ('E', 'Densidad', dens),
                                       ('F', 'Estado de la fibra', estado), ('G', 'Canas', canas)]:
        partes.append(
            f'IF(INDEX({SERV(col_serv)},MAX(1,{fila_serv}))="Sí",'
            f'SUMIFS({AJ(col_aj)},{AJ("A")},"{variable}",{AJ("B")},""&{valor}),0)'
        )
    return '+'.join(partes)


def precio(fila_serv, largo, dens, estado, canas):
    base = f'INDEX({SERV("C")},MAX(1,{fila_serv}))'
    redondeo = 'MAX(1,Servicios!$E$5)'
    sumas = suma_variables('D', fila_serv, largo, dens, estado, canas)
    pcts = suma_variables('E', fila_serv, largo, dens, estado, canas)
    return f'ROUND(({base}+{sumas})*(1+{pcts})/{redondeo},0)*{redondeo}'


# ══════════════════════════════════════════════════════════════════════════
# Simulador
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Simulador', 'Simulador de una cotización',
    'Elige el servicio y las variables. La hoja calcula el precio, el costo de producto, la comisión y el margen.',
    [38, 26, 50],
)
entradas_sim = [
    (4, 'Servicio', 'Ejemplo · Tinte completo', f'={SERV("A")}'),
    (5, 'Largo', 'Medio', f'={RANGO_LARGO}'),
    (6, 'Densidad', 'Alta', f'={RANGO_DENS}'),
    (7, 'Estado de la fibra', 'Sano', f'={RANGO_ESTADO}'),
    (8, 'Canas', '50–75 %', f'={RANGO_CANAS}'),
    (9, 'Estilista que atiende', 'Ejemplo · Caro', f'={RANGO_ESTILISTAS}'),
    (10, 'Tinte que se usa', 'Ejemplo · Tinte 60 g', f'={PROD("A")}'),
    (11, 'Oxidante que se usa', 'Ejemplo · Oxidante 20 vol', f'={PROD("A")}'),
]
ws['A3'] = 'Datos de la cotización'
ws['A3'].font = F_SUBTITULO
for fila, etiqueta, valor, fuente in entradas_sim:
    ws.cell(row=fila, column=1, value=etiqueta).font = F_ETIQUETA
    entrada(ws.cell(row=fila, column=2), valor)
    lista(ws, fuente, f'B{fila}')

calculos = [
    (13, 'sub', 'Precio', None, None),
    (14, 'f', 'Fila del servicio (dato interno)', f'=IFERROR(MATCH($B$4,{SERV("A")},0),0)', '0'),
    (15, 'f', 'Precio base', f'=IF($B$14=0,0,INDEX({SERV("C")},MAX(1,$B$14)))', COP),
    (16, 'f', 'Suma por variables', f'=IF($B$14=0,0,{suma_variables("D", "$B$14", "$B$5", "$B$6", "$B$7", "$B$8")})', COP),
    (17, 'f', 'Porcentaje por variables', f'=IF($B$14=0,0,{suma_variables("E", "$B$14", "$B$5", "$B$6", "$B$7", "$B$8")})', PCT),
    (18, 'f', 'Precio sin redondear', '=(B15+B16)*(1+B17)', COP),
    (19, 'total', 'Precio de la cotización', f'=IF($B$14=0,0,{precio("$B$14", "$B$5", "$B$6", "$B$7", "$B$8")})', COP),
    (20, 'f', 'Válida hasta', '=TODAY()+Servicios!$B$6', FECHA),
    (22, 'sub', 'Costo de producto (servicios con fórmula de color)', None, None),
    (23, 'f', '¿Usa fórmula de color?', f'=IF($B$14=0,"",INDEX({SERV("H")},MAX(1,$B$14)))', None),
    (24, 'f', 'Gramos de tinte base', f'=IF($B$14=0,0,INDEX({SERV("I")},MAX(1,$B$14)))', GRAMOS),
    (25, 'f', 'Factor por largo', f'=SUMIFS({CONS("C")},{CONS("A")},"Largo",{CONS("B")},$B$5)', FACTOR),
    (26, 'f', 'Factor por densidad', f'=SUMIFS({CONS("C")},{CONS("A")},"Densidad",{CONS("B")},$B$6)', FACTOR),
    (27, 'f', 'Gramos de tinte estimados', '=IF(B23="Sí",B24*B25*B26,0)', GRAMOS),
    (28, 'f', 'Costo del tinte', f'=B27*IFERROR(INDEX({PROD("G")},MATCH($B$10,{PROD("A")},0)),0)', COP),
    (29, 'f', 'Mililitros de oxidante', '=B27*Consumo!$B$4', GRAMOS),
    (30, 'f', 'Costo del oxidante', f'=B29*IFERROR(INDEX({PROD("G")},MATCH($B$11,{PROD("A")},0)),0)', COP),
    (31, 'f', 'Insumos fijos', f'=IF($B$14=0,0,IF(ISNUMBER(INDEX({SERV("J")},MAX(1,$B$14))),INDEX({SERV("J")},MAX(1,$B$14)),0))', COP),
    (32, 'total', 'Costo total de producto', '=B28+B30+B31', COP),
    (34, 'sub', 'Comisión y margen', None, None),
    (35, 'f', 'Fila de la estilista (dato interno)', f'=IFERROR(MATCH($B$9,{COM("A")},0),0)', '0'),
    (36, 'f', '% de comisión',
     f'=IF($B$35=0,0,IF(AND(OR(INDEX({SERV("B")},MAX(1,$B$14))="color",INDEX({SERV("B")},MAX(1,$B$14))="decoloración"),'
     f'ISNUMBER(INDEX({COM("C")},MAX(1,$B$35)))),INDEX({COM("C")},MAX(1,$B$35)),INDEX({COM("B")},MAX(1,$B$35))))', PCT),
    (37, 'f', 'Se calcula sobre', f'=IF($B$35=0,"Precio",IF(INDEX({COM("D")},MAX(1,$B$35))="","Precio",INDEX({COM("D")},MAX(1,$B$35))))', None),
    (38, 'f', 'Comisión', '=B36*IF(B37="Precio menos producto",B19-B32,B19)', COP),
    (39, 'total', 'Margen del salón', '=B19-B32-B38', COP),
    (40, 'f', 'Margen (%)', '=IF(B19>0,B39/B19,0)', PCT),
]
explicaciones = {
    16: 'Sumas fijas de Ajustes, solo de las variables que aplican a este servicio.',
    17: 'Porcentajes de Ajustes, solo de las variables que aplican a este servicio.',
    19: 'Redondeado al múltiplo de Servicios › Redondeo.',
    27: 'Gramos base × factor por largo × factor por densidad.',
    36: 'En color o decoloración usa el % de color si existe; si no, el general.',
    39: 'Precio − costo de producto − comisión.',
}
for fila, tipo, etiqueta, f, fmt in calculos:
    c = ws.cell(row=fila, column=1, value=etiqueta)
    if tipo == 'sub':
        c.font = F_SUBTITULO
        continue
    c.font = F_ETIQUETA if tipo == 'total' else F_BASE
    destino = ws.cell(row=fila, column=2)
    formula(destino, f, fmt, fuerte=tipo == 'total')
    if tipo == 'total':
        destino.fill = FILL_RESULTADO
    if fila in explicaciones:
        n = ws.cell(row=fila, column=3, value=explicaciones[fila])
        n.font = F_NOTA

# ══════════════════════════════════════════════════════════════════════════
# Registro
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Registro', 'Registro de la prueba (una semana)',
    'Una fila por cada servicio real. La hoja calcula su precio y tú escribes lo que se cobró. '
    'Solo iniciales de la clienta; nada de datos de salud capilar.',
    [12, 22, 12, 28, 12, 11, 15, 14, 15, 15, 13, 12, 32, 32, 9, 9],
)
encabezado(ws, 4, ['Fecha', 'Estilista', 'Clienta (iniciales)', 'Servicio', 'Largo', 'Densidad',
                   'Estado de la fibra', 'Canas', 'Precio según la hoja', 'Precio cobrado',
                   'Diferencia ($)', 'Diferencia (%)', '¿Por qué cambió?', '¿Algo que la hoja no considera?',
                   'Fila servicio', 'Dentro de ±5 %'])
REG_INI, REG_FIN = 5, 104
for i in range(REG_INI, REG_FIN + 1):
    ejemplo_fila = i == REG_INI
    valores = ([date.today(), 'Ejemplo · Caro', 'L. M.', 'Ejemplo · Tinte completo', 'Medio', 'Alta', 'Sano',
                '50–75 %'] if ejemplo_fila else [None] * 8)
    for col, v in enumerate(valores, start=1):
        entrada(ws.cell(row=i, column=col), v, None, ejemplo=ejemplo_fila)
    ws.cell(row=i, column=1).number_format = FECHA
    formula(ws.cell(row=i, column=15), f'=IF($D{i}="",0,IFERROR(MATCH($D{i},{SERV("A")},0),0))', '0')
    formula(ws.cell(row=i, column=9),
            f'=IF($O{i}=0,"",{precio(f"$O{i}", f"$E{i}", f"$F{i}", f"$G{i}", f"$H{i}")})', COP)
    entrada(ws.cell(row=i, column=10), 190000 if ejemplo_fila else None, COP, ejemplo=ejemplo_fila)
    formula(ws.cell(row=i, column=11), f'=IF(OR(I{i}="",J{i}=""),"",J{i}-I{i})', COP)
    formula(ws.cell(row=i, column=12), f'=IF(OR(I{i}="",J{i}="",I{i}=0),"",J{i}/I{i}-1)', PCT)
    entrada(ws.cell(row=i, column=13), 'Ejemplo: la clienta pidió matiz adicional' if ejemplo_fila else None,
            None, ejemplo=ejemplo_fila)
    entrada(ws.cell(row=i, column=14))
    formula(ws.cell(row=i, column=16), f'=IF(OR(I{i}="",J{i}=""),"",IF(ABS(J{i}-I{i})<=0.05*I{i},1,0))', '0')
lista(ws, f'={RANGO_ESTILISTAS}', f'B{REG_INI}:B{REG_FIN}')
lista(ws, f'={SERV("A")}', f'D{REG_INI}:D{REG_FIN}')
lista(ws, f'={RANGO_LARGO}', f'E{REG_INI}:E{REG_FIN}')
lista(ws, f'={RANGO_DENS}', f'F{REG_INI}:F{REG_FIN}')
lista(ws, f'={RANGO_ESTADO}', f'G{REG_INI}:G{REG_FIN}')
lista(ws, f'={RANGO_CANAS}', f'H{REG_INI}:H{REG_FIN}')
ws.column_dimensions.group('O', 'P', hidden=True)
ws.freeze_panes = f'E{REG_INI}'
REG = lambda col: f'Registro!${col}${REG_INI}:${col}${REG_FIN}'

# ══════════════════════════════════════════════════════════════════════════
# Resumen
# ══════════════════════════════════════════════════════════════════════════
ws = hoja(
    'Resumen', 'Resumen de la prueba',
    'Se llena solo con lo anotado en Registro. Responde: ¿el salón cobra lo que dice la hoja? ¿Quién se separa y cuánto?',
    [40, 14, 26, 22],
)
resumen = [
    (4, 'Casos con precio de la hoja y cobro', f'=COUNT({REG("P")})', '0'),
    (5, 'Cobrados al precio de la hoja (±5 %)', f'=IF(B4=0,0,SUM({REG("P")})/B4)', PCT),
    (6, 'Lo cobrado frente a la hoja (todos los casos)',
     f'=IF(SUMIFS({REG("I")},{REG("P")},">=0")=0,0,SUMIFS({REG("J")},{REG("P")},">=0")/SUMIFS({REG("I")},{REG("P")},">=0")-1)', PCT),
]
for fila, etiqueta, f, fmt in resumen:
    ws.cell(row=fila, column=1, value=etiqueta).font = F_ETIQUETA
    formula(ws.cell(row=fila, column=2), f, fmt, fuerte=True)
ws['C6'] = 'Positivo: se cobró más que la hoja. Negativo: menos.'
ws['C6'].font = F_NOTA
encabezado(ws, 8, ['Estilista', 'Casos', 'Cobró frente a la hoja', 'Casos fuera de ±5 %'])
for k in range(COM_FIN - COM_INI + 1):
    fila = 9 + k
    origen = COM_INI + k
    formula(ws.cell(row=fila, column=1), f'=IF(Comisiones!$A${origen}="","",Comisiones!$A${origen})', enlace=True)
    formula(ws.cell(row=fila, column=2),
            f'=IF(A{fila}="","",COUNTIFS({REG("B")},""&A{fila},{REG("P")},">=0"))', '0')
    formula(ws.cell(row=fila, column=3),
            f'=IF(A{fila}="","",IF(B{fila}=0,"",SUMIFS({REG("J")},{REG("B")},""&A{fila},{REG("P")},">=0")'
            f'/SUMIFS({REG("I")},{REG("B")},""&A{fila},{REG("P")},">=0")-1))', PCT)
    formula(ws.cell(row=fila, column=4),
            f'=IF(A{fila}="","",IF(B{fila}=0,"",COUNTIFS({REG("B")},""&A{fila},{REG("P")},0)))', '0')

# ── Fuente Arial en todas las celdas que quedaron con la fuente por defecto ─
for hoja_ws in wb.worksheets:
    for fila_celdas in hoja_ws.iter_rows():
        for c in fila_celdas:
            if c.font is None or c.font.name != 'Arial':
                c.font = F_BASE

wb.calculation.fullCalcOnLoad = True
wb.save(SALIDA)
print('ok', SALIDA)
